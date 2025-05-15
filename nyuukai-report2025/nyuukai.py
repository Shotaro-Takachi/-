# import bottle
# import sqlite3
# from datetime import datetime, timedelta
# from openpyxl import Workbook

# # データベース接続
# conn = sqlite3.connect("sqlite.db", check_same_thread=False)
# conn.execute("DROP TABLE IF EXISTS members")
# conn.execute("""
# CREATE TABLE IF NOT EXISTS members (
#     id INTEGER PRIMARY KEY AUTOINCREMENT,
#     name TEXT,
#     line_name TEXT,
#     date TEXT,
#     time TEXT,
#     method TEXT,
#     first_choice TEXT,
#     second_choice TEXT,
#     third_choice TEXT
# )
# """)
# conn.commit()

# # アプリケーション作成
# application = bottle.Bottle()

# # ルート（予約フォームの表示）
# @application.route("/")
# def root():
#     today = datetime.now().strftime("%Y-%m-%d")
#     max_date = (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d")
#     return bottle.template("./static/reservation.html", encoding='utf-8', today=today, max_date=max_date)

# # 提出処理（フォームからのデータを受け取る）
# @application.post("/submit")
# @bottle.view("submit")
# def submit():
#     bottle.response.content_type = "text/html; charset=utf-8"

#     # フォームデータの取得
#     name = bottle.request.forms.get("name")
#     line_name = bottle.request.forms.get("line_name")
#     date = bottle.request.forms.get("date")
#     time = bottle.request.forms.get("time")
#     method = bottle.request.forms.get("method")
#     first_choice = bottle.request.forms.get("first_choice")
#     second_choice = bottle.request.forms.get("second_choice")
#     third_choice = bottle.request.forms.get("third_choice")

#     # 日付のバリデーション
#     try:
#         date_obj = datetime.strptime(date, "%Y-%m-%d")
#         formatted_date = date_obj.strftime("%Y年%m月%d日")
#     except ValueError:
#         return bottle.HTTPError(400, "日付は「YYYY-MM-DD」の形式で入力してください。")

#     # 時間のバリデーション
#     try:
#         time_obj = datetime.strptime(time, "%H:%M")
#         formatted_time = time_obj.strftime("%H:%M")
#     except ValueError:
#         return bottle.HTTPError(400, "時間は「HH:MM」の形式で入力してください。")

#     # データベースに登録
#     sql_insert = """
#     INSERT INTO members (name, line_name, date, time, method, first_choice, second_choice, third_choice) 
#     VALUES (?, ?, ?, ?, ?, ?, ?, ?)
#     """
#     conn.execute(sql_insert, (
#         name, line_name, formatted_date, formatted_time, method, first_choice, second_choice, third_choice
#     ))
#     conn.commit()

#     # Excelファイルの作成
#     wb = Workbook()
#     ws = wb.active
#     ws.title = "予約一覧"
#     headers = ["ID", "氏名", "LINEの名前", "日付", "時間", "形式", "第一希望局", "第二希望局", "第三希望局"]
#     ws.append(headers)

#     # 登録データをExcelファイルに追加
#     ws.append([1, name, line_name, formatted_date, formatted_time, method, first_choice, second_choice, third_choice])
#     file_path = "nyuukai-form2025.xlsx"
#     wb.save(file_path)

#     # 登録内容を返す
#     return {
#         "name": name,
#         "line_name": line_name,
#         "date": formatted_date,
#         "time": formatted_time,
#         "method": method,
#         "first_choice": first_choice,
#         "second_choice": second_choice,
#         "third_choice": third_choice,
#         "file_path": file_path
#     }

# # アプリケーションの起動
# if __name__ == "__main__":
#      application.run(host='localhost', port=8080)

import bottle
import sqlite3
from datetime import datetime, timedelta
from openpyxl import Workbook

conn = sqlite3.connect("sqlite.db", check_same_thread=False)

conn.execute("DROP TABLE IF EXISTS members")
conn.execute("""
CREATE TABLE IF NOT EXISTS members
(id integer primary key autoincrement,
first_choice text, first_time text, second_choice text, second_time text, third_choice text,
name text, line_name text, date text, method text)
""")
conn.commit()

application = bottle.Bottle()

@application.route("/")
def root():
    today = datetime.now().strftime("%Y-%m-%d")  
    max_date = (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d")
    return bottle.template("./static/reservation.html", {"today": today, "max_date": max_date})

@application.post("/submit")
@bottle.view("submit")
def submit():
    first_choice = bottle.request.params.first_choice
    first_time = bottle.request.params.first_time
    second_choice = bottle.request.params.second_choice
    second_time = bottle.request.params.second_time
    third_choice = bottle.request.params.third_choice
    name = bottle.request.params.name
    line_name = bottle.request.params.line_name
    date = bottle.request.params.date
    method = bottle.request.params.method

    try:
        date_obj = datetime.strptime(date, "%Y-%m-%d")
        formatted_date = date_obj.strftime("%Y年%m月%d日")  
    except ValueError:
        return bottle.HTTPError(400, "日付は「YYYY-MM-DD」の形式で入力してください。")

    sql_check = """
    SELECT * FROM members
    WHERE date=? AND first_choice=? AND first_time=? 
    """
    cur = conn.cursor()
    cur.execute(sql_check, (date, first_choice, first_time))
    existing = cur.fetchone()

    if existing:
        return bottle.HTTPError(400, f"{first_choice}の{formatted_date}の面談予約は既に埋まっています。")
    
    sql_check = """
    SELECT * FROM members
    WHERE date=? AND second_choice=? AND second_time=? 
    """
    cur = conn.cursor()
    cur.execute(sql_check, (date, first_choice, first_time))
    existing = cur.fetchone()

    if existing:
        return bottle.HTTPError(400, f"{first_choice}の{formatted_date}の面談予約は既に埋まっています。")


    sql_insert = """
    INSERT INTO members (first_choice, first_time,second_choice, second_time,third_choice, name, line_name, date, method) 
    VALUES (?, ?, ?, ?, ?, ?, ?, ?,?)"""
    conn.execute(sql_insert, (first_choice, first_time,second_choice, second_time,third_choice, name, line_name, date,method))
    conn.commit()
    
    # Excelファイルの作成・更新
    wb = Workbook()
    ws = wb.active
    ws.title = "予約一覧"
    headers = ["ID", "第一希望", "第一希望時間", "第二希望", "第二希望時間", "第三希望", "氏名", "LINEの名前", "日付", "方法"]
    ws.append(headers)

    cur.execute("SELECT * FROM members ORDER BY date, first_time")
    rows = cur.fetchall()
    for row in rows:
        ws.append(row)

    file_path = "nyuukai-form2025.xlsx"
    wb.save(file_path)
    
    return {
        "name": name,
        "line_name": line_name,
        "date": formatted_date,
        "first_choice": first_choice,
        "first_time": first_time,
        "second_choice": second_choice,
        "second_time": second_time,
        "third_choice": third_choice,
        "method": method,
        "file_path": file_path
    }

# @application.route("/list")
# @bottle.view("list")
# def list():
#     cur = conn.cursor()
#     cur.execute("SELECT * FROM members ORDER BY date, time")
#     rows = cur.fetchall()
    
#     data = []
#     for row in rows:
#         try:
#             date_obj = datetime.strptime(row[6], "%Y-%m-%d") 
#             formatted_date = date_obj.strftime("%Y年%m月%d日") 
#         except ValueError:
#             formatted_date = row[6]
        
#         data.append({
#             "id": row[0],
#             "second_choice": row[2],
#             "third_choice": row[3],
#             "date": formatted_date,
#             "name": row[4],
#             "line_name": row[5],
#             "method": row[7],
#         })

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "予約一覧"

#     # ヘッダーの追加
#     headers = ["ID", "第二希望局", "第三希望局", "日程", "氏名", "ラインの名前", "方法"]
#     ws.append(headers)

#     # データの追加
#     for res in data:
#         ws.append([res["id"], res["second_choice"], res["third_choice"], res["date"], res["name"], res["line_name"], res["method"]])

#     # Excelファイルの保存
#     file_path = "nyuukai-form2025.xlsx"
#     wb.save(file_path)

#     return {"data": data, "file_path": file_path}

if __name__ == "__main__":
    application.run(host='localhost', port=8080)
